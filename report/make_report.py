"""Generate report/GreenSettle_report.docx (python-docx).

    python experiments/run_offline.py
    python report/make_figures.py
    python report/make_report.py

Every number comes from report/facts.py: results files written by scripts
in this repo, or the code's own config. Missing real-laptop results (E1)
appear as PENDING. Open the file in Word and accept "update fields" to fill
in the table of contents.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))

import facts  # noqa: E402

OUT = HERE / "GreenSettle_report.docx"


# ---------------------------------------------------------------------------
# python-docx helpers
# ---------------------------------------------------------------------------
def field(paragraph, instruction: str, placeholder: str = "") -> None:
    """Insert a Word field (TOC, PAGE, ...) into a paragraph."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, sep, text, end):
        run._r.append(element)


def update_fields_on_open(doc: Document) -> None:
    settings = doc.settings.element
    flag = OxmlElement("w:updateFields")
    flag.set(qn("w:val"), "true")
    settings.append(flag)


class Report:
    def __init__(self) -> None:
        self.doc = Document()
        self.figure_no = 0
        self.table_no = 0
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        section = self.doc.sections[0]
        section.left_margin = section.right_margin = Cm(2.3)
        section.top_margin = section.bottom_margin = Cm(2.2)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("Page ")
        field(footer, "PAGE", "1")

    def h(self, text: str, level: int = 1) -> None:
        self.doc.add_heading(text, level=level)

    def p(self, text: str, bold_prefix: str | None = None, italic: bool = False) -> None:
        para = self.doc.add_paragraph()
        if bold_prefix:
            para.add_run(bold_prefix).bold = True
        run = para.add_run(text)
        run.italic = italic

    def bullets(self, items: list[str]) -> None:
        for item in items:
            self.doc.add_paragraph(item, style="List Bullet")

    def numbered(self, items: list[str]) -> None:
        for item in items:
            self.doc.add_paragraph(item, style="List Number")

    def code(self, lines: list[str]) -> None:
        for line in lines:
            para = self.doc.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        self.doc.add_paragraph()

    def figure(self, file: str, caption: str, width_cm: float = 15.5) -> None:
        self.figure_no += 1
        self.doc.add_picture(str(facts.FIGURES / file), width=Cm(width_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.doc.add_paragraph(f"Figure {self.figure_no}: {caption}", style="Caption")

    def table(self, header: list[str], rows: list[list[str]], caption: str, widths: list[float] | None = None) -> None:
        self.table_no += 1
        self.doc.add_paragraph(f"Table {self.table_no}: {caption}", style="Caption")
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for cell, text in zip(t.rows[0].cells, header):
            cell.text = ""
            cell.paragraphs[0].add_run(text).bold = True
        for row in rows:
            cells = t.add_row().cells
            for cell, text in zip(cells, row):
                cell.text = text
        for row in t.rows:
            for i, cell in enumerate(row.cells):
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9.5)
                if widths:
                    cell.width = Cm(widths[i])
        self.doc.add_paragraph()

    def pending(self, text: str) -> None:
        run = self.doc.add_paragraph().add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xD0, 0x3B, 0x3B)


# ---------------------------------------------------------------------------
def build() -> Document:
    r = Report()
    head = facts.headlines()
    cap = facts.captions()

    # ---- title page -------------------------------------------------------------
    r.doc.add_paragraph("GreenSettle", style="Title")
    r.p("Carbon credits for cloud workloads, settled after the run from measured energy")
    r.p("TODO: author name · student number · course · date", italic=True)
    r.h("Contents", 1)
    field(r.doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u', "Open in Word and update fields to build the table of contents.")
    r.doc.add_section(WD_SECTION.NEW_PAGE)

    # ---- abstract -------------------------------------------------------------
    r.h("Abstract", 1)
    r.p("GreenSettle schedules batch jobs into greener time slots and settles carbon credits after each job has run, "
        "from the energy that was actually measured. A worker on a laptop measures the CPU package energy of each job "
        "with LibreHardwareMonitor and attributes it by CPU time, with an uncertainty estimate. The server compares "
        "the job with the same job started at submission time, using actual grid data, propagates the measurement "
        "and intensity uncertainties into the saving, and pays a capped rebate only for the part of the saving that "
        "lies beyond k standard uncertainties (a dead zone). The grid, the forecast and the noise models are "
        "simulated. " + head["e3"] + " " + head["e4"] + " " + head["gaming"] + " " + head["e1"])
    r.p(f"The contribution claimed is narrow: {facts.NOVELTY}. {facts.LIT_NOTE}")

    # ---- problem ---------------------------------------------------------------
    r.h("1. Problem", 1)
    r.p("TODO: quote the exact wording of Challenge 10 and Challenge 8 from the course brief here.", italic=True)
    r.p("Carbon-aware computing moves flexible work to hours when the electricity grid is cleaner. Incentives for "
        "such behaviour are usually granted from forecasts and estimates made before the work runs. Two questions "
        "are rarely answered afterwards: did the job really use the energy that was assumed, and was the moment it "
        "ran really cleaner than running it immediately? If rewards are paid on estimates, measurement noise and "
        "forecast errors can earn credits for savings that never happened.")
    r.p("GreenSettle addresses this with three parts: measurement of each job's energy with an uncertainty "
        "(Phase 1), an append-only credit ledger with a settlement rule that only pays for savings beyond the noise "
        "(Phase 2), and a scheduler, API and worker that put both into a running system (Phase 3).")

    # ---- related work ---------------------------------------------------------
    r.h("2. Related work (short)", 1)
    r.p(facts.LIT_NOTE + " The references below were not re-checked against their sources and are marked "
        "[UNVERIFIED]; they must be verified before submission.")
    r.bullets([
        "Temporal workload shifting in the cloud: Wiesner et al., \"Let's Wait Awhile: How Temporal Workload "
        "Shifting Can Reduce Carbon Emissions in the Cloud\", Middleware 2021. [UNVERIFIED]",
        "Carbon-aware scheduling at datacenter scale: Radovanovic et al., \"Carbon-Aware Computing for Datacenters\", "
        "IEEE Transactions on Power Systems. [UNVERIFIED]",
        "Per-software carbon accounting: Green Software Foundation, Software Carbon Intensity (SCI) specification. "
        "[UNVERIFIED]",
        "Energy measurement with Intel RAPL: Khan et al., \"RAPL in Action: Experiences in Using RAPL for Power "
        "Measurements\", ACM TOMPECS 2018. [UNVERIFIED]",
        "Uncertainty propagation: JCGM 100:2008, Guide to the expression of uncertainty in measurement (GUM). [UNVERIFIED]",
    ])
    r.p("TODO, papers still to read:", bold_prefix=None)
    r.bullets([
        "TODO: counterfactual baselines and additionality in carbon-offset markets.",
        "TODO: measurement and verification (M&V) protocols for energy savings (e.g. IPMVP).",
        "TODO: carbon-intensity forecasting accuracy (grid operators, Electricity Maps methodology).",
        "TODO: any prior work on settling scheduling incentives ex post with uncertainty (to test the novelty claim).",
    ])

    # ---- design -----------------------------------------------------------------
    r.h("3. Design", 1)
    r.figure("fig_a_architecture.png", cap["architecture"])
    r.h("3.1 Measurement (Phase 1)", 2)
    r.p("The worker polls the CPU Package power from LibreHardwareMonitor about once per second and integrates it "
        "with the trapezoidal rule. Idle power is calibrated separately and subtracted; the remaining dynamic energy "
        "is attributed to the job by its share of busy CPU time (psutil, including child processes). The "
        "uncertainty combines sensor tolerance, sampling error, idle-baseline noise and attribution error in "
        "quadrature. The scope is the CPU package only.")
    r.h("3.2 Settlement (Phase 2)", 2)
    r.p("compute_settlement (greensettle/settlement/rules.py) takes the job energy E ± dE, the counterfactual "
        "intensity i_submit (the same job started at submission, from actual grid data) and the intensity during "
        "the run i_run, each with an uncertainty:")
    r.code([
        "fee        = p · E · i_run                          (actual emissions)",
        "saved      = E · (i_submit − i_run)",
        "saved_unc  = √((dE·(i_submit − i_run))² + (E·u_submit)² + (E·u_run)²)",
        "saving     if saved − k·saved_unc > 0:  rebate = min(ρ·p·(saved − k·saved_unc), γ·fee)",
        "backfire   elif −saved − k·saved_unc > 0: surcharge = σ·p·(−saved − k·saved_unc)",
        "dead zone  otherwise: no rebate, no surcharge",
        "net        = fee + surcharge − rebate",
    ])
    r.p("Because the rebate is at most γ·fee with γ < 1, a job with a positive fee always has a positive net cost, "
        "so running jobs cannot earn credits. Rebates are paid on the pessimistic saving, so a noise-sized saving "
        "earns nothing.")
    r.figure("fig_b_deadzone.png", cap["deadzone"])
    r.p("The ledger is append-only (SQLite triggers block UPDATE and DELETE), stores integer micro-credits, rounds "
        "every amount against the user, applies a lazy half-life decay before each balance-changing write, and "
        "derives a tier (red, standard, green) from the balance after decay.")
    r.h("3.3 Scheduling (Phase 3)", 2)
    r.p("choose_slot (greensettle/scheduler/slots.py) plans each job once at submission. It asks the forecast for "
        "every 30-minute start inside the user's wait window, picks the lowest forecast (earliest on ties), and "
        "defers only if the forecast saving exceeds defer_z combined forecast uncertainties; otherwise the job "
        "starts now. Ready jobs are dispatched by overdue first (a starvation guard), then tier priority, planned "
        "start and submission order, within global and per-tier concurrency limits. The server runs on a simulated "
        "clock so that a day of grid variation fits in a demonstration.")

    # ---- implementation ---------------------------------------------------------
    r.h("4. Implementation", 1)
    r.bullets([
        "Python 3.11+, standard library plus psutil; FastAPI and uvicorn for the server; SQLite with versioned "
        "migrations (PRAGMA user_version).",
        "greensettle/agent: LibreHardwareMonitor client, meter, worker. greensettle/ledger: migrations, ledger, "
        "decay, allowance, settle_job, CLI. greensettle/scheduler: slots, planner, dispatch, queue. "
        "greensettle/server: API, clock, dashboard.",
        "The worker pulls work (POST /agent/claim) so the laptop can sit behind NAT; claims are atomic "
        "(BEGIN IMMEDIATE) and leased; an expired lease requeues a job once.",
        "A single-file dashboard (no external resources) shows the simulated grid, live power, the queue with "
        "candidate slots, settlements and balances.",
        "The test suite runs without hardware (mock sensor, fake clock); hardware tests are opt-in.",
    ])

    # ---- security ---------------------------------------------------------------
    r.h("5. Security model", 1)
    r.bullets([
        "Jobs are never commands: {job_type, params} from a fixed catalog with integer parameters in hard ranges. "
        "The worker re-validates and builds a fixed argv itself; no shell, no user strings reach a process; burn "
        "jobs are capped at 60 s for thermal safety.",
        "Authentication: per-user API keys (32 random bytes, shown once, stored as SHA-256); agent and admin "
        "tokens only from environment variables; constant-time comparison.",
        "Limits: queued jobs per user, request body size, strict request models that reject unknown fields, wrong "
        "types and NaN; error responses do not echo input.",
        "Network: binds to 127.0.0.1 unless told otherwise; remote use needs HTTPS or an SSH tunnel.",
        "Trust model: the server trusts the worker's energy numbers. Measurement attestation is out of scope; a "
        "dishonest worker could lie.",
    ])

    # ---- experiments --------------------------------------------------------------
    r.h("6. Experiments and results", 1)
    r.p("E1 runs on the real laptop. E3, E4 and the gaming table are simulations: the grid curve, the forecast "
        "error and the measurement noise are simulated, and are labelled so below.")
    r.h("6.1 E1: measurement repeatability (real laptop)", 2)
    d1 = facts.e1()
    if d1:
        r.table(["workers", "runs", "mean E_job (J)", "sd (J)", "CV", "coverage ±1σ"],
                [[str(s["workers"]), str(s["runs"]), f"{s['mean_kwh'] * 3.6e6:.1f}", f"{s['sd_kwh'] * 3.6e6:.1f}",
                  f"{s['cv']:.1%}", f"{s['coverage_1sigma']:.0%}"] for s in d1["summary"]],
                f"E1 on {d1['machine']['platform']} ({d1['created_at']}), burn {d1['settings']['seconds']} s.")
    else:
        r.pending(f"E1: {facts.PENDING}. Command: python experiments/e1_repeatability.py --repeats 6 --seconds 20 --workers 1,4")
    r.h("6.2 E3: scheduling benefit (simulated)", 2)
    r.figure("fig_c_e3.png", cap["e3"])
    d3 = facts.e3()
    rows = [x for x in d3["rows"] if x["delay_h"] > 0 and x["forecast_error"] == 0.1]
    r.table(["delay (h)", "defer_z", "change vs immediate", "mean wait (h)", "deferred", "deferrals backfired"],
            [[str(x["delay_h"]), f"{x['defer_z']:g}", f"{x['change_pct']:+.1f} %", f"{x['mean_wait_h']:.2f}",
              f"{x['share_deferred']:.0%}", f"{x['share_deferrals_backfired']:.0%}"] for x in rows],
            "E3 at 10 % simulated forecast error (full grid of settings in experiments/results/e3.csv).")
    r.p(head["e3"])
    r.h("6.3 E4: dead-zone width (simulated)", 2)
    r.figure("fig_d_e4.png", cap["e4"])
    d4 = facts.e4()
    by = {(x["k"], x["true_gain"]): x for x in d4["rows"]}
    ks = sorted({x["k"] for x in d4["rows"]})
    r.table(["k", "false-reward rate (gain 0)", "reward rate (gain 20 %)", "reward rate (gain 40 %)",
             "credits minted, gain 0"],
            [[f"{k:g}", f"{by[(k, 0.0)]['reward_rate']:.1%}", f"{by[(k, 0.2)]['reward_rate']:.1%}",
              f"{by[(k, 0.4)]['reward_rate']:.1%}", f"{by[(k, 0.0)]['credits_minted']:.1f}"] for k in ks],
            "E4: share of jobs that received a rebate, against k (simulated noise model).")
    r.p(head["e4"])
    r.h("6.4 Gaming table (simulated)", 2)
    g = facts.gaming()
    r.table(["archetype", "rule", "mean net / job", "share rebated", "jobs with net < 0"],
            [[x["archetype"], x["rule"], f"{float(x['mean_net_credits']):.2f}", f"{float(x['share_rebated']):.1%}",
              x["negative_net_jobs"]] for x in g],
            "Gaming table: 200 simulated jobs per archetype; the ex-ante rule is a simplified strawman.")
    r.p(head["gaming"])

    # ---- limitations ----------------------------------------------------------------
    r.h("7. Limitations", 1)
    r.bullets([
        "CPU package energy only: RAM, display, storage and network are not measured; the laptop's memory and "
        "platform power sensors read 0 W.",
        "Attribution by CPU time ignores that P-core and E-core seconds cost different energy on this hybrid CPU.",
        "The server trusts worker measurements; there is no attestation.",
        "Jobs are planned once and never re-planned; one mean intensity is used per job.",
        "Failed jobs are not settled, so their energy is never billed.",
        "Overdue jobs skip tier priority; a job submitted with no wait window is overdue almost immediately.",
        "One laptop, one simulated grid, simulated forecast errors and noise; all parameters are placeholders "
        "(Table below). Only a limited literature search was done.",
    ])

    # ---- placeholders ----------------------------------------------------------------
    r.h("8. Placeholder parameters", 1)
    r.p("None of these values was calibrated. Each is read from the code that uses it.")
    r.table(["parameter", "value", "defined in"], [[n, v, w] for n, v, w in facts.placeholders()],
            "Placeholder parameters (values read from the code).", widths=[4.2, 7.0, 5.2])

    # ---- conclusion -----------------------------------------------------------------
    r.h("9. Conclusion", 1)
    r.p("GreenSettle shows that carbon incentives for scheduled jobs can be settled after the run, from measured "
        "energy and actual grid data, with uncertainty carried through to the payment. In the simulations, "
        "scheduling within a wait window lowered emissions per job, the dead zone suppressed rewards for savings "
        "that were only noise, and no ex-post settlement ever paid a job to run. The real-hardware repeatability "
        "of the measurement (E1) and the calibration of every placeholder are the next steps.")
    if not d1:
        r.pending(f"E1 is {facts.PENDING}; this conclusion does not yet include real-laptop evidence.")

    # ---- appendix -------------------------------------------------------------------
    r.h("Appendix A: Reproduce", 1)
    r.code([
        "python -m venv .venv && .venv\\Scripts\\Activate.ps1",
        "pip install -r requirements.txt -r requirements-report.txt",
        "python -m pytest                                  # hardware tests skipped",
        "python experiments/run_offline.py                 # worked examples, gaming table, E3, E4, SUMMARY.md",
        "python -m greensettle.cli calibrate-idle --seconds 30          # on the laptop",
        "python experiments/e1_repeatability.py --repeats 6 --seconds 20 --workers 1,4",
        "python report/make_figures.py",
        "python report/make_poster.py",
        "python report/make_report.py",
    ])
    update_fields_on_open(r.doc)
    return r.doc


def main() -> int:
    for needed in ("e3.json", "e4.json", "gaming_table.csv", "worked_examples.json"):
        if not (facts.RESULTS / needed).exists():
            sys.exit(f"missing experiments/results/{needed}: run python experiments/run_offline.py")
    if not (facts.FIGURES / "captions.json").exists():
        sys.exit("missing figures: run python report/make_figures.py")
    build().save(OUT)
    print(f"Wrote {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
